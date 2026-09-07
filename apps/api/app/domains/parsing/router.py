import io
from typing import Optional
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
import openpyxl
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.core.models import BoQ, Document, LineItem, utc_now
from app.core.security import AuthContext, require_contractor
from app.domains.audit.service import log_audit_event
from app.integrations.llm import get_llm_provider
from app.integrations.storage import get_storage_provider
from app.schemas.boq import BoQResponse, ParseRequest

router = APIRouter(tags=["Parsing"])


def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        extracted = []
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted.append(f"--- PAGE {page_idx + 1} ---\n{text}")
        return "\n\n".join(extracted)
    except Exception:
        return ""


def extract_text_from_excel(xlsx_bytes: bytes) -> str:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"--- SHEET: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
                if row_vals:
                    lines.append(" | ".join(row_vals))
        return "\n".join(lines)
    except Exception:
        return ""


@router.post("/boqs/{boq_id}/parse", response_model=BoQResponse)
async def parse_boq_document(
    boq_id: str,
    payload: Optional[ParseRequest] = None,
    auth: AuthContext = Depends(require_contractor),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(BoQ)
        .options(selectinload(BoQ.source_document), selectinload(BoQ.line_items))
        .where(BoQ.id == boq_id, BoQ.contractor_organisation_id == auth.organisation_id)
    )
    res = await db.execute(stmt)
    boq = res.scalar_one_or_none()
    if not boq:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={"code": "BOQ_NOT_FOUND", "message": "BoQ not found."},
        )

    llm_provider = get_llm_provider()
    try:
        if payload and payload.pasted_text and payload.pasted_text.strip():
            parse_result = await llm_provider.parse_boq_document(
                extracted_text=payload.pasted_text.strip(),
                filename=None,
                context={"region": boq.region},
            )
        elif boq.source_document:
            doc = boq.source_document
            storage = get_storage_provider()
            doc_bytes = await storage.get_file(doc.storage_key)
            parse_result = await llm_provider.parse_boq_document_bytes(
                document_bytes=doc_bytes,
                mime_type=doc.mime_type,
                filename=doc.original_filename,
                context={"region": boq.region},
            )
        else:
            parse_result = await llm_provider.parse_boq_document(
                extracted_text="",
                filename=None,
                context={"region": boq.region},
            )
    except ValueError as ve:
        raw_err = str(ve)
        if "SCANNED_PDF_NO_TEXT_LAYER" in raw_err:
            err_code = "SCANNED_PDF_NO_TEXT_LAYER"
            clean_message = (
                "We couldn't read this document properly — it appears to be a scanned document or image-only PDF without a searchable digital text layer. "
                "Try re-exporting it as a text-based PDF, or upload an Excel version if you have one."
            )
        elif "GARBLED_DOCUMENT_TEXT" in raw_err:
            err_code = "GARBLED_DOCUMENT_TEXT"
            sub_reason = raw_err.replace("GARBLED_DOCUMENT_TEXT:", "").strip()
            clean_message = (
                f"We couldn't read this document properly ({sub_reason}) — "
                "try re-exporting it as a text-based PDF, or upload an Excel version if you have one."
            )
        elif "BINARY_FILE" in raw_err or "TEXT_DECODING_FAILED" in raw_err:
            err_code = "BINARY_FILE"
            clean_message = (
                "We couldn't read this file — it contains unreadable binary data. "
                "Please upload a standard Excel (.xlsx) or text-based PDF document."
            )
        else:
            err_code = "GARBLED_DOCUMENT_TEXT"
            clean_message = (
                f"We couldn't read this document properly: {raw_err} — "
                "try re-exporting it as a text-based PDF, or upload an Excel version if you have one."
            )

        await log_audit_event(
            db=db,
            organisation_id=auth.organisation_id,
            actor_user_id=auth.user_id,
            entity_type="boq",
            entity_id=boq.id,
            action="boq.parse_failed_corrupted_text",
            metadata_json={
                "error": raw_err,
                "error_code": err_code,
                "filename": boq.source_document.original_filename if boq.source_document else None,
            },
            after_json={
                "error": raw_err,
                "error_code": err_code,
                "filename": boq.source_document.original_filename if boq.source_document else None,
            },
        )
        await db.commit()

        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"code": err_code, "message": clean_message},
        )

    if not parse_result.line_items:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"code": "PARSING_FAILED", "message": "Could not extract any valid line items from the document."},
        )

    if boq.title == "New Tender BoQ" and parse_result.title_hint:
        boq.title = parse_result.title_hint
    if not boq.tender_reference and parse_result.tender_reference_hint:
        boq.tender_reference = parse_result.tender_reference_hint

    # Clear previous items
    boq.line_items.clear()
    await db.flush()

    for item_dto in parse_result.line_items:
        item = LineItem(
            boq_id=boq.id,
            source_row_reference=item_dto.source_row_reference,
            description=item_dto.description,
            unit=item_dto.unit,
            quantity=item_dto.quantity,
            category=item_dto.category,
            benchmark_min_minor=item_dto.benchmark_min_minor,
            benchmark_max_minor=item_dto.benchmark_max_minor,
            benchmark_source=item_dto.benchmark_source,
            benchmark_currency="ZAR",
            pricing_status="unsourced",
            parsing_confidence=item_dto.parsing_confidence,
            review_status=item_dto.review_status,
            exclusion_reason=item_dto.exclusion_reason,
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        boq.line_items.append(item)

    # Also store excluded candidates so contractor can inspect/restore them if desired
    for item_dto in parse_result.excluded_candidates:
        item = LineItem(
            boq_id=boq.id,
            source_row_reference=item_dto.source_row_reference,
            description=item_dto.description,
            unit=item_dto.unit,
            quantity=item_dto.quantity,
            category=item_dto.category,
            benchmark_min_minor=None,
            benchmark_max_minor=None,
            benchmark_source=None,
            benchmark_currency="ZAR",
            pricing_status="unsourced",
            parsing_confidence=item_dto.parsing_confidence,
            review_status="excluded",
            exclusion_reason=item_dto.exclusion_reason or "Low confidence / filtered candidate",
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        boq.line_items.append(item)

    boq.status = "parsed"
    boq.updated_at = utc_now()

    valid_count = len(parse_result.line_items)
    excluded_count = len(parse_result.excluded_candidates)
    needs_review_count = sum(1 for i in parse_result.line_items if i.review_status == "needs_review")
    high_conf_count = sum(1 for i in parse_result.line_items if i.review_status == "accepted")

    await log_audit_event(
        db=db,
        organisation_id=auth.organisation_id,
        actor_user_id=auth.user_id,
        entity_type="boq",
        entity_id=boq.id,
        action="boq.parsed",
        after_json={
            "valid_items_count": valid_count,
            "excluded_items_count": excluded_count,
            "needs_review_count": needs_review_count,
            "high_confidence_count": high_conf_count,
            "detected_title": parse_result.title_hint,
            "detected_ref": parse_result.tender_reference_hint,
            "metadata": parse_result.metadata,
        },
    )

    await db.commit()
    await db.refresh(boq)
    return BoQResponse.model_validate(boq)


@router.post("/boqs/{boq_id}/parse-file", response_model=BoQResponse)
async def parse_boq_file(
    boq_id: str,
    file: UploadFile = File(...),
    auth: AuthContext = Depends(require_contractor),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(BoQ)
        .options(selectinload(BoQ.source_document), selectinload(BoQ.line_items))
        .where(BoQ.id == boq_id, BoQ.contractor_organisation_id == auth.organisation_id)
    )
    res = await db.execute(stmt)
    boq = res.scalar_one_or_none()
    if not boq:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={"code": "BOQ_NOT_FOUND", "message": "BoQ not found."},
        )

    doc_bytes = await file.read()
    if not doc_bytes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"code": "EMPTY_FILE", "message": "Uploaded file is empty."},
        )

    filename = file.filename or "uploaded_document"
    mime_type = file.content_type or "application/octet-stream"

    # Pre-check for binary files disguised as text/csv
    is_text_ext = filename.endswith(".csv") or filename.endswith(".txt")
    if (is_text_ext or "text" in mime_type) and (doc_bytes.startswith(b"PK\x03\x04") or b"\x00" in doc_bytes[:1024]):
        err_code = "BINARY_FILE"
        clean_message = (
            "We couldn't read this file — it contains unreadable binary data. "
            "Please upload a standard Excel (.xlsx) or text-based PDF document."
        )
        await log_audit_event(
            db=db,
            organisation_id=auth.organisation_id,
            actor_user_id=auth.user_id,
            entity_type="boq",
            entity_id=boq.id,
            action="boq.parse_failed_corrupted_text",
            metadata_json={"error_code": err_code, "error": "BINARY_FILE", "filename": filename},
            after_json={"error_code": err_code, "error": "BINARY_FILE", "filename": filename},
        )
        await db.commit()
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"code": err_code, "message": clean_message},
        )

    llm_provider = get_llm_provider()
    try:
        parse_result = await llm_provider.parse_boq_document_bytes(
            document_bytes=doc_bytes,
            mime_type=mime_type,
            filename=filename,
            context={"region": boq.region},
        )
    except ValueError as ve:
        raw_err = str(ve)
        if "SCANNED_PDF_NO_TEXT_LAYER" in raw_err:
            err_code = "SCANNED_PDF_NO_TEXT_LAYER"
            clean_message = (
                "We couldn't read this document properly — it appears to be a scanned document or image-only PDF without a searchable digital text layer. "
                "Try re-exporting it as a text-based PDF, or upload an Excel version if you have one."
            )
        elif "GARBLED_DOCUMENT_TEXT" in raw_err:
            err_code = "GARBLED_DOCUMENT_TEXT"
            sub_reason = raw_err.replace("GARBLED_DOCUMENT_TEXT:", "").strip()
            clean_message = (
                f"We couldn't read this document properly ({sub_reason}) — "
                "try re-exporting it as a text-based PDF, or upload an Excel version if you have one."
            )
        elif "BINARY_FILE" in raw_err or "TEXT_DECODING_FAILED" in raw_err:
            err_code = "BINARY_FILE"
            clean_message = (
                "We couldn't read this file — it contains unreadable binary data. "
                "Please upload a standard Excel (.xlsx) or text-based PDF document."
            )
        else:
            err_code = "GARBLED_DOCUMENT_TEXT"
            clean_message = (
                f"We couldn't read this document properly: {raw_err} — "
                "try re-exporting it as a text-based PDF, or upload an Excel version if you have one."
            )

        await log_audit_event(
            db=db,
            organisation_id=auth.organisation_id,
            actor_user_id=auth.user_id,
            entity_type="boq",
            entity_id=boq.id,
            action="boq.parse_failed_corrupted_text",
            metadata_json={"error_code": err_code, "error": raw_err, "filename": filename},
            after_json={"error_code": err_code, "error": raw_err, "filename": filename},
        )
        await db.commit()

        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"code": err_code, "message": clean_message},
        )

    if not parse_result.line_items:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"code": "PARSING_FAILED", "message": "Could not extract any valid line items from the document."},
        )

    if boq.title == "New Tender BoQ" and parse_result.title_hint:
        boq.title = parse_result.title_hint
    if not boq.tender_reference and parse_result.tender_reference_hint:
        boq.tender_reference = parse_result.tender_reference_hint

    boq.line_items.clear()
    await db.flush()

    for item_dto in parse_result.line_items:
        item = LineItem(
            boq_id=boq.id,
            source_row_reference=item_dto.source_row_reference,
            description=item_dto.description,
            unit=item_dto.unit,
            quantity=item_dto.quantity,
            category=item_dto.category,
            benchmark_min_minor=item_dto.benchmark_min_minor,
            benchmark_max_minor=item_dto.benchmark_max_minor,
            benchmark_source=item_dto.benchmark_source,
            benchmark_currency="ZAR",
            pricing_status="unsourced",
            parsing_confidence=item_dto.parsing_confidence,
            review_status=item_dto.review_status,
            exclusion_reason=item_dto.exclusion_reason,
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        boq.line_items.append(item)

    for item_dto in parse_result.excluded_candidates:
        item = LineItem(
            boq_id=boq.id,
            source_row_reference=item_dto.source_row_reference,
            description=item_dto.description,
            unit=item_dto.unit,
            quantity=item_dto.quantity,
            category=item_dto.category,
            benchmark_min_minor=None,
            benchmark_max_minor=None,
            benchmark_source=None,
            benchmark_currency="ZAR",
            pricing_status="unsourced",
            parsing_confidence=item_dto.parsing_confidence,
            review_status="excluded",
            exclusion_reason=item_dto.exclusion_reason or "Low confidence / filtered candidate",
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        boq.line_items.append(item)

    boq.status = "parsed"
    boq.updated_at = utc_now()

    await log_audit_event(
        db=db,
        organisation_id=auth.organisation_id,
        actor_user_id=auth.user_id,
        entity_type="boq",
        entity_id=boq.id,
        action="boq.parsed",
        after_json={
            "valid_items_count": len(parse_result.line_items),
            "excluded_items_count": len(parse_result.excluded_candidates),
            "detected_title": parse_result.title_hint,
            "detected_ref": parse_result.tender_reference_hint,
        },
    )

    await db.commit()
    await db.refresh(boq)
    return BoQResponse.model_validate(boq)

