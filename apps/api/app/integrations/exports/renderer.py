import io
from datetime import datetime
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


class StandardExportRenderer:
    async def render_excel(
        self,
        boq_data: Dict[str, Any],
        line_items_data: List[Dict[str, Any]],
        audit_events: List[Dict[str, Any]],
    ) -> bytes:
        wb = openpyxl.Workbook()
        ws_boq = wb.active
        ws_boq.title = "Priced BoQ Schedule"

        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Arial", size=10, italic=True, color="4B5563")
        bold_font = Font(name="Arial", size=10, bold=True)
        regular_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )
        currency_format = 'R #,##0.00'

        # Document Header
        ws_boq["A1"] = boq_data.get("title", "Tender Bill of Quantities")
        ws_boq["A1"].font = title_font
        ws_boq["A2"] = f"Tender Ref: {boq_data.get('tender_reference', 'N/A')} | Region: {boq_data.get('region', 'N/A')} | Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws_boq["A2"].font = sub_font

        # Table Headers
        headers = [
            "Item", "Ref", "Description", "Category", "Unit", "Qty",
            "Unit Rate (ZAR)", "Total (ZAR)", "Status", "Supplier / Quote Ref"
        ]
        start_row = 4
        for col_idx, header in enumerate(headers, 1):
            cell = ws_boq.cell(row=start_row, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 5, 6, 9] else "left", vertical="center")

        # Table Rows
        total_minor = 0
        current_row = start_row + 1
        for idx, item in enumerate(line_items_data, 1):
            qty = float(item.get("quantity", 1.0))
            price_minor = item.get("final_price_minor")
            unit_rate = (price_minor / 100.0) if price_minor is not None else 0.0
            line_total = unit_rate * qty
            if price_minor:
                total_minor += int(price_minor * qty)

            status = item.get("pricing_status", "unsourced").replace("_", " ").title()
            quote_ref = item.get("quote_reference") or item.get("override_reason") or "-"

            row_values = [
                idx,
                item.get("source_row_reference") or f"Item-{idx}",
                item.get("description", ""),
                item.get("category", "General").title(),
                item.get("unit", "no"),
                qty,
                unit_rate,
                line_total,
                status,
                quote_ref,
            ]

            for col_idx, val in enumerate(row_values, 1):
                cell = ws_boq.cell(row=current_row, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in [1, 2, 5]:
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in [6]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = '#,##0.00'
                elif col_idx in [7, 8]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = currency_format

            current_row += 1

        # Total Row
        total_row = current_row
        ws_boq.cell(row=total_row, column=7, value="TOTAL TENDER PRICE:").font = bold_font
        ws_boq.cell(row=total_row, column=7).alignment = Alignment(horizontal="right")
        total_cell = ws_boq.cell(row=total_row, column=8, value=total_minor / 100.0)
        total_cell.font = bold_font
        total_cell.number_format = currency_format
        total_cell.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
        total_cell.border = thin_border

        # Adjust column widths
        for col in ws_boq.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_boq.column_dimensions[col_letter].width = max(max_len + 3, 10)
        ws_boq.column_dimensions['C'].width = 40

        # Sheet 2: Audit Trail
        ws_audit = wb.create_sheet(title="Audit Trail")
        ws_audit["A1"] = "Procurement Pricing & Override Audit Log"
        ws_audit["A1"].font = title_font

        audit_headers = ["Timestamp", "Entity", "Action", "Actor", "Before", "After", "Details"]
        for col_idx, h in enumerate(audit_headers, 1):
            cell = ws_audit.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font

        for idx, event in enumerate(audit_events, 4):
            ws_audit.cell(row=idx, column=1, value=str(event.get("created_at", ""))).font = regular_font
            ws_audit.cell(row=idx, column=2, value=str(event.get("entity_type", "")).title()).font = regular_font
            ws_audit.cell(row=idx, column=3, value=str(event.get("action", ""))).font = regular_font
            ws_audit.cell(row=idx, column=4, value=str(event.get("actor_name") or event.get("actor_user_id") or "System")).font = regular_font
            ws_audit.cell(row=idx, column=5, value=str(event.get("before_json") or "-")).font = regular_font
            ws_audit.cell(row=idx, column=6, value=str(event.get("after_json") or "-")).font = regular_font
            ws_audit.cell(row=idx, column=7, value=str(event.get("metadata_json") or "-")).font = regular_font

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def render_pdf(
        self,
        boq_data: Dict[str, Any],
        line_items_data: List[Dict[str, Any]],
        audit_events: List[Dict[str, Any]],
    ) -> bytes:
        if not REPORTLAB_AVAILABLE:
            # Fallback text representation if reportlab is unavailable
            content = f"TENDERPRENEUR BOQ EXPORT\nTitle: {boq_data.get('title')}\n"
            return content.encode('utf-8')

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1E3A8A'),
            spaceAfter=6,
        )
        meta_style = ParagraphStyle(
            'MetaStyle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#4B5563'),
            spaceAfter=12,
        )
        table_text = ParagraphStyle('TableText', parent=styles['Normal'], fontSize=8)
        table_bold = ParagraphStyle('TableBold', parent=styles['Normal'], fontSize=8, fontName='Helvetica-Bold')

        # Header
        elements.append(Paragraph(f"<b>Tenderpreneur Priced Bill of Quantities</b>: {boq_data.get('title', 'Tender')}", title_style))
        elements.append(Paragraph(f"Tender Reference: <b>{boq_data.get('tender_reference', 'N/A')}</b> | Region: <b>{boq_data.get('region', 'N/A')}</b> | Date: {datetime.now().strftime('%Y-%m-%d')}", meta_style))
        elements.append(Spacer(1, 10))

        # Line Items Table
        table_data = [
            ["#", "Ref", "Description", "Unit", "Qty", "Rate (ZAR)", "Total (ZAR)", "Pricing Status & Evidence"]
        ]

        total_minor = 0
        for idx, item in enumerate(line_items_data, 1):
            qty = float(item.get("quantity", 1.0))
            price_minor = item.get("final_price_minor")
            unit_rate = (price_minor / 100.0) if price_minor is not None else 0.0
            line_total = unit_rate * qty
            if price_minor:
                total_minor += int(price_minor * qty)

            status = item.get("pricing_status", "unsourced").replace("_", " ").title()
            quote_ref = item.get("quote_reference") or item.get("override_reason") or "-"
            evidence = f"{status}\n({quote_ref})"

            table_data.append([
                str(idx),
                str(item.get("source_row_reference") or f"Item-{idx}"),
                Paragraph(item.get("description", ""), table_text),
                item.get("unit", "no"),
                f"{qty:,.2f}",
                f"R {unit_rate:,.2f}",
                f"R {line_total:,.2f}",
                Paragraph(evidence, table_text),
            ])

        # Summary Row
        table_data.append([
            "", "", Paragraph("<b>TOTAL TENDER AMOUNT</b>", table_bold), "", "", "",
            Paragraph(f"<b>R {total_minor / 100.0:,.2f}</b>", table_bold), ""
        ])

        t = Table(table_data, colWidths=[25, 45, 260, 40, 50, 80, 90, 160])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ALIGN', (0, 0), (1, -1), 'CENTER'),
            ('ALIGN', (3, 0), (3, -1), 'CENTER'),
            ('ALIGN', (4, 0), (6, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FEF08A')),
        ]))
        elements.append(t)
        doc.build(elements)
        return buffer.getvalue()
